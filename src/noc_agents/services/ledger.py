"""Excel shift-ledger writer.

Split in two since the transactional outbox (spec §7.0.2 / §5.3.9):

* ``ledger_row_cells`` is pure: it renders the file name and the row cells from the
  incident. The LEDGER node calls it inside the transaction and queues the result as an
  ``EXCEL_ROW`` outbox row.
* ``append_excel_row`` does the file I/O under a file lock. Only the outbox dispatcher
  calls it, after the incident has committed.

``write_excel_row`` composes the two and is kept for callers that want the old
one-call behaviour (nothing in the hot path uses it any more).
"""

from __future__ import annotations

import os
import threading
import time
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook

from noc_agents.config import OperatorConfig
from noc_agents.db.models import IncidentRow

ROOT = Path(__file__).resolve().parents[3]

HEADERS = [
    "Time (EAT)",
    "Incident No",
    "Priority",
    "Site ID",
    "Site Name",
    "Type",
    "Class",
    "Region",
    "TT Category",
    "Failure Domain",
    "Est. Users",
    "Responsible MSP",
    "FE",
    "RNIO",
    "Escalated At",
    "Expected Resolve",
    "Status",
    "Vendor TT",
    "M-PESA Risk",
    "Shift",
]

LOCK_TIMEOUT_S = 10.0  # how long an append waits for another writer's lock file
LOCK_STALE_S = 60.0  # a lock file older than this belongs to a crashed writer


def ledger_root() -> Path:
    """Shift-ledger folder: LEDGER_DIR env var, else the default ROOT/data/shift_ledgers."""
    raw = (os.getenv("LEDGER_DIR") or "").strip()
    return Path(raw) if raw else ROOT / "data" / "shift_ledgers"


def eat_now(cfg: OperatorConfig) -> datetime:
    return datetime.now(ZoneInfo(cfg.timezone))


def ledger_file_name(shift_type: str, when: datetime) -> str:
    return f"ledger_{when.strftime('%Y-%m-%d')}_{shift_type.upper()}.xlsx"


def ledger_row_cells(inc: IncidentRow, cfg: OperatorConfig, shift_type: str) -> tuple[str, list]:
    """Render (file name, row cells) for the incident. Pure: no file is touched."""
    eat = eat_now(cfg)
    cells = [
        eat.strftime("%Y-%m-%d %H:%M:%S"),
        inc.incident_number,
        inc.priority,
        inc.site_id,
        inc.site_name,
        inc.site_type,
        getattr(inc, "site_class", "") or "",
        inc.region_code,
        getattr(inc, "tt_category", "") or "",
        inc.failure_domain,
        inc.users_affected,
        getattr(inc, "responsible_msp", None) or inc.msp_name or "",
        inc.fe_name or "",
        getattr(inc, "rnio_name", "") or "",
        str(getattr(inc, "escalated_at", "") or ""),
        str(getattr(inc, "expected_resolution_at", "") or ""),
        inc.status,
        getattr(inc, "vendor_tt_ref", "") or "",
        "YES" if inc.mpesa_risk else "NO",
        shift_type.upper(),
    ]
    return ledger_file_name(shift_type, eat), cells


def append_excel_row(operator_id: str, file_name: str, cells: list) -> Path:
    """Append one row to ``<LEDGER_DIR>/<operator>/<file_name>`` under a file lock.

    Raises OSError (PermissionError when the workbook is open in Excel, TimeoutError when
    another writer holds the lock) so the outbox dispatcher can classify it as transient
    and retry. Called only from the dispatcher, never inside a transaction.
    """
    if Path(file_name).name != file_name:  # the payload came from the DB: never let it choose a folder
        raise ValueError(f"ledger file name must be a bare name, got {file_name!r}")
    folder = ledger_root() / operator_id
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / file_name
    with _file_lock(path):
        if path.exists():
            wb = load_workbook(path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "ShiftFailures"
            ws.append(HEADERS)
        ws.append(cells)
        wb.save(path)
    return path


def write_excel_row(inc: IncidentRow, cfg: OperatorConfig, shift_type: str) -> Path:
    """Render and append in one call (compatibility; the hot path enqueues instead)."""
    file_name, cells = ledger_row_cells(inc, cfg, shift_type)
    return append_excel_row(cfg.operator_id, file_name, cells)


# --- file lock ----------------------------------------------------------------------------

_thread_locks: dict[str, threading.Lock] = {}
_thread_locks_guard = threading.Lock()


def _thread_lock(path: Path) -> threading.Lock:
    with _thread_locks_guard:
        return _thread_locks.setdefault(str(path), threading.Lock())


@contextmanager
def _file_lock(path: Path):
    """Exclusive ``<path>.lock`` (O_EXCL) plus an in-process lock, so two dispatcher
    threads or two processes never interleave a load/append/save on the same workbook."""
    lock_path = path.with_name(path.name + ".lock")
    with _thread_lock(path):
        deadline = time.monotonic() + LOCK_TIMEOUT_S
        while True:
            try:
                fd = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
                break
            except FileExistsError:
                try:
                    if time.time() - lock_path.stat().st_mtime > LOCK_STALE_S:
                        lock_path.unlink(missing_ok=True)  # crashed writer: take over
                        continue
                except OSError:
                    pass  # vanished between exists and stat: loop and retry
                if time.monotonic() > deadline:
                    raise TimeoutError(f"shift ledger lock busy: {lock_path}")
                time.sleep(0.05)
        try:
            os.close(fd)
            yield
        finally:
            lock_path.unlink(missing_ok=True)
